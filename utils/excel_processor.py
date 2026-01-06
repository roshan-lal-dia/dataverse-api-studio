"""
Excel and CSV file processor
Handles reading Excel/CSV files, detecting headers, and extracting data
"""

import csv
import pandas as pd
from pathlib import Path
from typing import List, Dict, Tuple, Optional, Any


class ExcelProcessor:
    """Process Excel and CSV files for data import"""
    
    def __init__(self, file_path: str):
        """Initialize with file path"""
        self.file_path = Path(file_path)
        self.file_extension = self.file_path.suffix.lower()
        
        if self.file_extension not in ['.xlsx', '.xls', '.csv']:
            raise ValueError(f"Unsupported file format: {self.file_extension}")
    
    def read_file(self, header_row: Optional[int] = None) -> Tuple[List[str], List[Dict]]:
        """
        Read file and return headers + data rows
        Args:
            header_row: Override header row (0-indexed), None for auto-detect
        Returns:
            (headers, rows) where rows is list of dicts
        """
        if self.file_extension == '.csv':
            return self._read_csv(header_row)
        else:
            return self._read_excel(header_row)
    
    def _read_csv(self, header_row: Optional[int] = None) -> Tuple[List[str], List[Dict]]:
        """Read CSV file"""
        try:
            # Try pandas first for better handling
            if header_row is not None:
                df = pd.read_csv(self.file_path, header=header_row)
            else:
                # Auto-detect header
                detected_header = self._detect_header_row_csv()
                df = pd.read_csv(self.file_path, header=detected_header)
            
            # Clean column names (strip whitespace)
            df.columns = [str(col).strip() for col in df.columns]
            
            # Convert to list of dicts
            headers = list(df.columns)
            rows = df.to_dict('records')
            
            return headers, rows
        
        except Exception as e:
            raise Exception(f"Error reading CSV file: {str(e)}")
    
    def _read_excel(self, header_row: Optional[int] = None) -> Tuple[List[str], List[Dict]]:
        """Read Excel file with enhanced handling for merged cells and edge cases"""
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            ws = wb.active
            
            # Detect header row if not specified
            if header_row is None:
                header_row = self._detect_header_row_excel(ws)
            
            # Read headers with improved handling
            headers = []
            header_row_cells = list(ws[header_row + 1])  # openpyxl is 1-indexed
            
            for cell in header_row_cells:
                value = cell.value
                
                # Handle merged cells - get value from merged range if needed
                if value is None and cell.coordinate in ws.merged_cells:
                    # Find the merged range and get its value
                    for merged_range in ws.merged_cells.ranges:
                        if cell.coordinate in merged_range:
                            # Get value from top-left cell of merged range
                            top_left = merged_range.start_cell
                            value = top_left.value
                            break
                
                # Convert value to string and handle various types
                if value is not None:
                    # Handle numeric headers (convert to string)
                    if isinstance(value, (int, float)):
                        headers.append(f"Col_{int(value)}")
                    else:
                        headers.append(str(value).strip())
                else:
                    # Generate fallback column name
                    headers.append(f"Column_{cell.column_letter}")
            
            # Ensure unique headers
            headers = self._ensure_unique_headers(headers)
            
            # Read data rows
            rows = []
            for row_idx, row_cells in enumerate(ws.iter_rows(min_row=header_row + 2)):
                row_dict = {}
                for col_idx, cell in enumerate(row_cells):
                    if col_idx < len(headers):
                        value = cell.value
                        
                        # Handle merged cells in data rows
                        if value is None and cell.coordinate in ws.merged_cells:
                            for merged_range in ws.merged_cells.ranges:
                                if cell.coordinate in merged_range:
                                    value = merged_range.start_cell.value
                                    break
                        
                        # Convert value to appropriate type
                        row_dict[headers[col_idx]] = self._normalize_value(value)
                
                # Skip completely empty rows
                if any(v for v in row_dict.values() if v != "" and v is not None):
                    rows.append(row_dict)
            
            wb.close()
            return headers, rows
        
        except Exception as e:
            raise Exception(f"Error reading Excel file: {str(e)}")
    
    def _detect_header_row_csv(self, max_rows: int = 5) -> int:
        """
        Auto-detect header row in CSV (scan first 5 rows)
        Returns 0-indexed row number
        """
        try:
            with open(self.file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                rows = [next(reader, None) for _ in range(max_rows)]
            
            return self._analyze_header_candidates(rows)
        
        except Exception:
            return 0  # Default to first row
    
    def _detect_header_row_excel(self, worksheet, max_rows: int = 5) -> int:
        """
        Auto-detect header row in Excel (scan first 5 rows)
        Returns 0-indexed row number
        """
        try:
            rows = []
            for row_idx in range(1, max_rows + 1):
                row_values = [cell.value for cell in worksheet[row_idx]]
                rows.append(row_values)
            
            return self._analyze_header_candidates(rows)
        
        except Exception:
            return 0  # Default to first row
    
    def _analyze_header_candidates(self, rows: List[List]) -> int:
        """
        Analyze candidate rows to determine which is the header
        Returns 0-indexed row number
        """
        scores = []
        
        for row in rows:
            if not row or len(row) == 0:
                scores.append(0)
                continue
            
            score = 0
            
            # Check for text-like content
            text_count = sum(1 for val in row if isinstance(val, str) and val.strip())
            
            # Check for non-numeric content
            non_numeric_count = sum(1 for val in row if not self._is_numeric(val))
            
            # Check for unique values (headers are usually unique)
            # Guard against empty row to prevent division by zero
            unique_ratio = len(set(str(v) for v in row if v)) / len(row) if len(row) > 0 else 0
            
            # Calculate score
            score = text_count * 2 + non_numeric_count + (unique_ratio * 10)
            scores.append(score)
        
        # Return row with highest score
        if scores:
            return scores.index(max(scores))
        
        return 0
    
    def _is_numeric(self, value: Any) -> bool:
        """Check if value is numeric"""
        if value is None or value == "":
            return False
        
        try:
            float(value)
            return True
        except (ValueError, TypeError):
            return False
    
    def _ensure_unique_headers(self, headers: List[str]) -> List[str]:
        """Ensure all headers are unique by appending suffixes to duplicates"""
        seen = {}
        unique_headers = []
        
        for header in headers:
            if header not in seen:
                seen[header] = 0
                unique_headers.append(header)
            else:
                seen[header] += 1
                unique_headers.append(f"{header}_{seen[header]}")
        
        return unique_headers
    
    def _normalize_value(self, value: Any) -> Any:
        """Normalize cell value to appropriate Python type"""
        if value is None:
            return ""
        
        # Handle dates and times
        from datetime import datetime, date, time
        if isinstance(value, (datetime, date, time)):
            return str(value)
        
        # Keep numbers as-is
        if isinstance(value, (int, float)):
            return value
        
        # Convert strings and strip whitespace
        if isinstance(value, str):
            return value.strip()
        
        # Convert other types to string
        return str(value)
    
    def get_preview_data(self, num_rows: int = 5) -> Dict:
        """
        Get preview of first few rows with headers
        Returns: {"headers": [...], "preview": [...]}
        """
        try:
            headers, rows = self.read_file()
            
            preview_rows = rows[:num_rows]
            
            return {
                "headers": headers,
                "preview": preview_rows,
                "total_rows": len(rows)
            }
        
        except Exception as e:
            return {"error": str(e)}
