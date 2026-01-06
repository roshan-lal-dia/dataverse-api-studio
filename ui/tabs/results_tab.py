"""
Results Tab - Enhanced with relationship navigation and Power BI export
"""

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QTextEdit, QTableWidget, QTableWidgetItem, QGroupBox,
    QTabWidget, QFileDialog, QMessageBox, QLabel, QMenu
)
from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QAction
import json
import csv
from typing import Optional
from datetime import datetime


class ResultsTab(QWidget):
    """Results display tab with enhanced export and navigation"""
    
    navigate_to_related = pyqtSignal(str, str, str)  # entity, field, record_id
    
    def __init__(self):
        super().__init__()
        
        self.current_result = None
        self.client = None
        self.navigation_history = []
        
        self._create_ui()
    
    def _create_ui(self):
        """Create UI components"""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)
        
        # Navigation breadcrumb (for relationship navigation)
        self.breadcrumb_container = QWidget()
        breadcrumb_layout = QHBoxLayout(self.breadcrumb_container)
        breadcrumb_layout.setContentsMargins(0, 0, 0, 0)
        
        self.breadcrumb_label = QLabel("📍 View: Main Results")
        self.breadcrumb_label.setStyleSheet("font-weight: bold;")
        breadcrumb_layout.addWidget(self.breadcrumb_label)
        
        self.back_button = QPushButton("⬅️ Back")
        self.back_button.clicked.connect(self._navigate_back)
        self.back_button.setVisible(False)
        breadcrumb_layout.addWidget(self.back_button)
        
        breadcrumb_layout.addStretch()
        layout.addWidget(self.breadcrumb_container)
        
        # Result display tabs
        self.result_tabs = QTabWidget()
        
        # JSON view
        json_widget = QWidget()
        json_layout = QVBoxLayout(json_widget)
        
        self.json_display = QTextEdit()
        self.json_display.setReadOnly(True)
        self.json_display.setPlaceholderText("Results will appear here...")
        json_layout.addWidget(self.json_display)
        
        # Table view with context menu
        table_widget = QWidget()
        table_layout = QVBoxLayout(table_widget)
        
        self.table_display = QTableWidget()
        self.table_display.setAlternatingRowColors(True)
        self.table_display.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table_display.customContextMenuRequested.connect(self._show_context_menu)
        self.table_display.cellDoubleClicked.connect(self._on_cell_double_click)
        table_layout.addWidget(self.table_display)
        
        self.result_tabs.addTab(json_widget, "JSON View")
        self.result_tabs.addTab(table_widget, "Table View")
        
        layout.addWidget(self.result_tabs)
        
        # Action buttons
        button_layout = QHBoxLayout()
        
        self.export_json_button = QPushButton("💾 Export JSON")
        self.export_json_button.clicked.connect(self._export_json)
        button_layout.addWidget(self.export_json_button)
        
        self.export_csv_button = QPushButton("💾 Export CSV")
        self.export_csv_button.clicked.connect(self._export_csv)
        button_layout.addWidget(self.export_csv_button)
        
        self.export_excel_button = QPushButton("📊 Export Excel")
        self.export_excel_button.clicked.connect(self._export_excel)
        button_layout.addWidget(self.export_excel_button)
        
        self.export_powerbi_button = QPushButton("📈 Power BI Export")
        self.export_powerbi_button.clicked.connect(self._export_powerbi)
        button_layout.addWidget(self.export_powerbi_button)
        
        self.copy_button = QPushButton("📋 Copy to Clipboard")
        self.copy_button.clicked.connect(self._copy_to_clipboard)
        button_layout.addWidget(self.copy_button)
        
        button_layout.addStretch()
        
        layout.addLayout(button_layout)
    
    def display_result(self, operation_data: dict):
        """Display operation result"""
        self.current_result = operation_data
        
        # Display in JSON view
        json_text = json.dumps(operation_data, indent=2)
        self.json_display.setPlainText(json_text)
        
        # Display in table view if applicable
        result = operation_data.get("result", {})
        
        if "data" in result:
            data = result["data"]
            
            # Handle list of records (query results)
            if isinstance(data, list) and data:
                self._populate_table(data)
            # Handle single record
            elif isinstance(data, dict):
                self._populate_table([data])
    
    def _populate_table(self, records: list):
        """Populate table widget with records"""
        if not records:
            return
        
        # Get all unique keys
        all_keys = set()
        for record in records:
            all_keys.update(record.keys())
        
        headers = sorted(list(all_keys))
        
        # Setup table
        self.table_display.setRowCount(len(records))
        self.table_display.setColumnCount(len(headers))
        self.table_display.setHorizontalHeaderLabels(headers)
        
        # Populate data
        for row_idx, record in enumerate(records):
            for col_idx, header in enumerate(headers):
                value = record.get(header, "")
                
                # Convert to string for display
                if isinstance(value, (dict, list)):
                    value_str = json.dumps(value)
                else:
                    value_str = str(value) if value is not None else ""
                
                item = QTableWidgetItem(value_str)
                self.table_display.setItem(row_idx, col_idx, item)
        
        # Resize columns
        self.table_display.resizeColumnsToContents()
    
    def _export_json(self):
        """Export results as JSON"""
        if not self.current_result:
            QMessageBox.warning(self, "No Data", "No results to export.")
            return
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export JSON",
            "results.json",
            "JSON Files (*.json)"
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(self.current_result, f, indent=2)
                
                QMessageBox.information(self, "Exported", f"Results exported to {file_path}")
            
            except Exception as e:
                QMessageBox.critical(self, "Export Error", f"Failed to export:\n{str(e)}")
    
    def _export_csv(self):
        """Export results as CSV"""
        if not self.current_result:
            QMessageBox.warning(self, "No Data", "No results to export.")
            return
        
        # Extract data
        result = self.current_result.get("result", {})
        data = result.get("data", [])
        
        if not data:
            QMessageBox.warning(self, "No Data", "No tabular data to export.")
            return
        
        if not isinstance(data, list):
            data = [data]
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export CSV",
            "results.csv",
            "CSV Files (*.csv)"
        )
        
        if file_path:
            try:
                # Get all keys
                all_keys = set()
                for record in data:
                    all_keys.update(record.keys())
                
                headers = sorted(list(all_keys))
                
                with open(file_path, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.DictWriter(f, fieldnames=headers)
                    writer.writeheader()
                    
                    for record in data:
                        # Convert complex types to strings
                        row = {}
                        for key in headers:
                            value = record.get(key, "")
                            if isinstance(value, (dict, list)):
                                row[key] = json.dumps(value)
                            else:
                                row[key] = value
                        
                        writer.writerow(row)
                
                QMessageBox.information(self, "Exported", f"Results exported to {file_path}")
            
            except Exception as e:
                QMessageBox.critical(self, "Export Error", f"Failed to export:\n{str(e)}")
    
    def _copy_to_clipboard(self):
        """Copy JSON results to clipboard"""
        from PyQt6.QtWidgets import QApplication
        
        if not self.current_result:
            QMessageBox.warning(self, "No Data", "No results to copy.")
            return
        
        json_text = json.dumps(self.current_result, indent=2)
        
        clipboard = QApplication.clipboard()
        clipboard.setText(json_text)
        
        QMessageBox.information(self, "Copied", "Results copied to clipboard.")
    
    def _export_excel(self):
        """Export results as Excel with formatting"""
        if not self.current_result:
            QMessageBox.warning(self, "No Data", "No results to export.")
            return
        
        result = self.current_result.get("result", {})
        data = result.get("data", [])
        
        if not data:
            QMessageBox.warning(self, "No Data", "No tabular data to export.")
            return
        
        if not isinstance(data, list):
            data = [data]
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Excel",
            "results.xlsx",
            "Excel Files (*.xlsx)"
        )
        
        if file_path:
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
                
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Results"
                
                # Get headers
                all_keys = set()
                for record in data:
                    all_keys.update(record.keys())
                headers = sorted(list(all_keys))
                
                # Write headers with formatting
                for col_idx, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")
                
                # Write data
                for row_idx, record in enumerate(data, start=2):
                    for col_idx, header in enumerate(headers, start=1):
                        value = record.get(header, "")
                        
                        # Convert complex types
                        if isinstance(value, (dict, list)):
                            value = json.dumps(value)
                        
                        ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Auto-size columns
                for col in ws.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column].width = adjusted_width
                
                wb.save(file_path)
                
                QMessageBox.information(self, "Exported", f"Results exported to {file_path}")
            
            except Exception as e:
                QMessageBox.critical(self, "Export Error", f"Failed to export:\n{str(e)}")
    
    def _export_powerbi(self):
        """Export results as Power BI-ready CSV (uppercase headers, stable types)"""
        if not self.current_result:
            QMessageBox.warning(self, "No Data", "No results to export.")
            return
        
        result = self.current_result.get("result", {})
        data = result.get("data", [])
        
        if not data:
            QMessageBox.warning(self, "No Data", "No tabular data to export.")
            return
        
        if not isinstance(data, list):
            data = [data]
        
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Power BI CSV",
            "results_powerbi.csv",
            "CSV Files (*.csv)"
        )
        
        if file_path:
            try:
                # Get headers
                all_keys = set()
                for record in data:
                    all_keys.update(record.keys())
                
                # Convert to uppercase for Power BI
                headers = sorted(list(all_keys))
                headers_upper = [h.upper() for h in headers]
                
                with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:  # UTF-8 with BOM for Excel
                    writer = csv.writer(f)
                    writer.writerow(headers_upper)
                    
                    for record in data:
                        row = []
                        for key in headers:
                            value = record.get(key, "")
                            
                            # Normalize types for Power BI
                            if isinstance(value, bool):
                                row.append("TRUE" if value else "FALSE")
                            elif isinstance(value, (dict, list)):
                                row.append(json.dumps(value))
                            elif value is None:
                                row.append("")
                            else:
                                row.append(str(value))
                        
                        writer.writerow(row)
                
                QMessageBox.information(
                    self, 
                    "Exported", 
                    f"Power BI-ready CSV exported to {file_path}\n\n"
                    "Features:\n"
                    "• Uppercase headers\n"
                    "• UTF-8 with BOM encoding\n"
                    "• Normalized data types"
                )
            
            except Exception as e:
                QMessageBox.critical(self, "Export Error", f"Failed to export:\n{str(e)}")
    
    def _show_context_menu(self, position):
        """Show context menu on table right-click"""
        if not self.table_display.selectedItems():
            return
        
        menu = QMenu(self)
        
        # Add navigate to related option
        navigate_action = QAction("🔗 Navigate to Related Records", self)
        navigate_action.triggered.connect(self._navigate_to_related_from_menu)
        menu.addAction(navigate_action)
        
        menu.addSeparator()
        
        # Copy cell value
        copy_action = QAction("📋 Copy Cell Value", self)
        copy_action.triggered.connect(self._copy_cell_value)
        menu.addAction(copy_action)
        
        menu.exec(self.table_display.viewport().mapToGlobal(position))
    
    def _copy_cell_value(self):
        """Copy selected cell value"""
        from PyQt6.QtWidgets import QApplication
        
        selected = self.table_display.selectedItems()
        if selected:
            value = selected[0].text()
            clipboard = QApplication.clipboard()
            clipboard.setText(value)
    
    def _on_cell_double_click(self, row: int, column: int):
        """Handle cell double-click - show full value in dialog"""
        item = self.table_display.item(row, column)
        if item:
            value = item.text()
            
            if len(value) > 100:  # Show dialog for long values
                msg = QMessageBox(self)
                msg.setWindowTitle("Cell Value")
                msg.setText(f"Column: {self.table_display.horizontalHeaderItem(column).text()}")
                msg.setDetailedText(value)
                msg.exec()
    
    def _navigate_to_related_from_menu(self):
        """Navigate to related records from context menu"""
        # TODO: Implement relationship navigation
        QMessageBox.information(
            self,
            "Coming Soon",
            "Relationship navigation will be implemented in the next phase!"
        )
    
    def _navigate_back(self):
        """Navigate back in navigation history"""
        if self.navigation_history:
            previous_result = self.navigation_history.pop()
            self.display_result(previous_result)
            
            if not self.navigation_history:
                self.back_button.setVisible(False)
                self.breadcrumb_label.setText("📍 View: Main Results")
    
    def set_client(self, client):
        """Set the Dataverse client"""
        self.client = client
    
    def export_results(self):
        """Export results (called from menu)"""
        self._export_json()
